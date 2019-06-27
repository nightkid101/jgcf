import re
import os
import subprocess
from docx import Document

from xlrd import open_workbook, xldate_as_tuple
from utility import cn2dig, get_year, request_site_page, get_content_text, remove_strip, format_date
from bs4 import BeautifulSoup as bs
from oss_utils import oss_add_file, init_ali_oss
from urllib.parse import urljoin

ali_bucket = init_ali_oss()


# parse table content
def parse_table(file_id, content_table_list, origin_title, origin_publish_date, db, logger):
    strip_list = ['', '\xa0', '-', '——', '———', '—', '/', '---', '----', '--', '-----', '------']

    announcement_code = person_name = person_company = company_name = legal_representative \
        = truth = punishment_basement = punishment_decision = organization = publish_date \
        = litigant = ''
    for each_row in content_table_list:
        if re.search(r'行政处罚决定书?文书?号', remove_strip(each_row[0].strip())):
            announcement_code = each_row[-1].strip()
        if re.search(r'(被处罚当事人)', remove_strip(each_row[0].strip())) and \
                re.search(r'(个人)', remove_strip(each_row[1].strip())) and \
                re.search(r'(姓名|名称)', remove_strip(each_row[2].strip())):
            person_name = each_row[-1].strip()
        if re.search(r'(被处罚当事人)', remove_strip(each_row[0].strip())) and \
                re.search(r'(个人)', remove_strip(each_row[1].strip())) and \
                re.search(r'(单位)', remove_strip(each_row[2].strip())):
            person_company = each_row[-1].strip()
        if re.search(r'(被处罚当事人)', remove_strip(each_row[0].strip())) and \
                re.search(r'(单位)', remove_strip(each_row[1].strip())) and \
                re.search(r'(名称)', remove_strip(each_row[2].strip())):
            company_name = each_row[-1].strip()
        if re.search(r'(被处罚当事人)', remove_strip(each_row[0].strip())) and \
                re.search(r'(单位)', remove_strip(each_row[1].strip())) and \
                re.search(r'(法定代表人|主要负责人)', remove_strip(each_row[2].strip())):
            legal_representative = each_row[-1].strip()

        if re.search(r'(主要(违法)?违规事实|案由)', remove_strip(each_row[0].strip())):
            truth = each_row[-1].strip()
        if re.search(r'(^行政处罚依据$)', remove_strip(each_row[0].strip())):
            punishment_basement = each_row[-1].strip()
        if re.search(r'(^行政处罚决定$)', remove_strip(each_row[0].strip())):
            punishment_decision = each_row[-1].strip()
        if re.search(r'(作出.*?机关名称)', remove_strip(each_row[0].strip())):
            organization = each_row[-1].strip()
        if re.search(r'(作出.*?日期)', remove_strip(each_row[0].strip())):
            publish_date = each_row[-1].strip()

    if person_name not in strip_list:
        litigant += '个人姓名: ' + remove_strip(person_name) + '\n'
    if person_company not in strip_list:
        litigant += '个人就职单位: ' + remove_strip(person_company) + '\n'
    if company_name not in strip_list:
        litigant += '单位名称: ' + remove_strip(company_name) + '\n'
    if legal_representative not in strip_list:
        litigant += '单位法定代表人姓名: ' + remove_strip(legal_representative)

    if publish_date != '':
        publish_date = format_date(publish_date)
    else:
        publish_date = format_date(origin_publish_date)

    title = remove_strip(origin_title)

    result_map = {
        'announcementTitle': title,
        'announcementOrg': '广西银保监局',
        'announcementDate': publish_date,
        'announcementCode': remove_strip(announcement_code),
        'facts': truth,
        'defenseOpinion': '',
        'defenseResponse': '',
        'litigant': litigant.strip(),
        'punishmentBasement': '',
        'punishmentDecision': ('依据' + punishment_basement + '，' + punishment_decision).replace('。，', '，').replace(
            '依据根据', '依据').replace('依据依据', '依据'),
        'type': '行政处罚决定',
        'oss_file_id': file_id,
        'status': 'not checked'
    }
    logger.info(result_map)
    if db.announcement.find({'announcementTitle': title, 'oss_file_id': file_id}).count() == 0:
        db.announcement.insert_one(result_map)
        logger.info('广西保监局 数据解析 ' + organization + ' -- 数据导入完成')
    else:
        logger.info('广西保监局 数据解析 ' + organization + ' -- 数据已经存在')
    db.parsed_data.update_one({'_id': file_id}, {'$set': {'parsed': True}})
    logger.info('广西保监局 数据解析 ' + organization + ' -- 修改parsed完成')


def guangxi_circ(db, logger):
    for each_circ_data in db.circ_data.find({'origin': '广西保监局', 'status': {'$nin': ['ignored']}}):
        announcement_url = each_circ_data['url']
        announcement_title = each_circ_data['title']

        if db.circ_data.find(
                {'url': announcement_url, 'status': 'parsed'}).count() == 1 and db.parsed_data.find(
            {'origin_url': announcement_url, 'parsed': True}).count() == 1:
            continue

        logger.info('广西保监局 ' + 'Url to parse: %s' % announcement_url)

        r = request_site_page(announcement_url)
        if r is None:
            logger.error('网页请求错误 %s' % announcement_url)
            continue
        content_soup = bs(r.content, 'lxml') if r else bs('', 'lxml')

        table_content = content_soup.find(id='tab_content')
        if not table_content:
            logger.error('网页请求错误 %s' % announcement_url)
            continue
        content_text = get_content_text(table_content.find_all('tr')[3])
        if content_text == '':
            continue
        title = table_content.find_all('tr')[0].text.strip()

        if '行政处罚信息主动公开事项' in title or '行政处罚主动公开事项' in title:
            if len([each_link.attrs['href'] for each_link in table_content.find_all('a')
                    if '.xls' in each_link.attrs['href'].strip()]) > 0:
                xlsx_link = [urljoin(announcement_url, each_link.attrs['href'])
                             for each_link in table_content.find_all('a')
                             if '.xls' in each_link.attrs['href'].strip()][0]

                response = request_site_page(xlsx_link)
                link_type = xlsx_link.split('.')[-1]
                with open('./test/tmp.' + link_type, 'wb') as f:
                    f.write(response.content)

                with open('./test/tmp.' + link_type, 'rb') as xlsx_file:
                    xlsx_content = xlsx_file.read()

                if db.parsed_data.find(
                        {'origin_url': announcement_url, 'oss_file_origin_url': xlsx_link}).count() == 0:
                    oss_file_map = {
                        'origin_url': announcement_url,
                        'oss_file_origin_url': xlsx_link,
                        'origin_url_id': each_circ_data['_id'],
                        'oss_file_type': 'xlsx',
                        'oss_file_name': announcement_title,
                        'oss_file_content': xlsx_content,
                        'parsed': False
                    }
                    insert_response = db.parsed_data.insert_one(oss_file_map)
                    file_id = insert_response.inserted_id
                    oss_add_file(ali_bucket, str(file_id) + '/' + announcement_title + '.xlsx', xlsx_content)
                    db.circ_data.update_one({'_id': each_circ_data['_id']}, {'$set': {'status': 'parsed'}})
                else:
                    db.circ_data.update_one({'_id': each_circ_data['_id']}, {'$set': {'status': 'parsed'}})
                    file_id = db.parsed_data.find_one({'origin_url': announcement_url,
                                                       'oss_file_origin_url': xlsx_link})['_id']

                excel_data = open_workbook('./test/tmp.' + link_type)
                logger.info('删除tmp文件')
                if os.path.exists('./test/tmp.xls'):
                    os.remove('./test/tmp.xls')
                if os.path.exists('./test/tmp.xlsx'):
                    os.remove('./test/tmp.xlsx')

                sheet = excel_data.sheets()[0]

                result_map_list = []
                for i in range(sheet.nrows):
                    if i < 2:
                        continue
                    document_code = sheet.cell(i, 1).value.strip()
                    if document_code == '':
                        document_code = document_code
                    real_title = '广西银保监局行政处罚决定书' + '(' + str(document_code) + ')'
                    litigant = sheet.cell(i, 2).value.strip()
                    if sheet.cell(i, 6).ctype == 3:
                        publish_date = xldate_as_tuple(sheet.cell_value(i, 6),
                                                       excel_data.datemode)
                        publish_date = str(publish_date[0]) + '年' + str(
                            publish_date[1]) + '月' + str(
                            publish_date[2]) + '日'
                    else:
                        try:
                            publish_date = str(sheet.cell(i, 6).value.year) + '年' + str(
                                sheet.cell(i, 6).value.month) + '月' + str(
                                sheet.cell(i, 6).value.day) + '日'
                        except:
                            publish_date = publish_date
                    truth = litigant + sheet.cell(i, 3).value
                    defense = defense_response = ''
                    punishment_decision = \
                        sheet.cell(i, 4).value + '，我局对' + litigant + '作出以下处罚：' + sheet.cell(i, 5).value
                    punishment_basis = ''
                    each_map = {
                        'announcementTitle': real_title,
                        'announcementOrg': '广西银保监局',
                        'announcementDate': publish_date,
                        'announcementCode': document_code,
                        'facts': truth,
                        'defenseOpinion': defense,
                        'defenseResponse': defense_response,
                        'litigant': litigant[:-1] if litigant[-1] == '：' else litigant,
                        'punishmentBasement': punishment_basis,
                        'punishmentDecision': punishment_decision,
                        'type': '行政处罚决定',
                        'oss_file_id': file_id,
                        'status': 'not checked'
                    }
                    logger.info(each_map)
                    if db.announcement.find(
                            {'announcementTitle': real_title,
                             'oss_file_id': file_id,
                             'litigant': each_map['litigant']}).count() == 0:
                        db.announcement.insert_one(each_map)
                        logger.info('广西保监局 数据解析 ' + ' -- 数据导入完成')
                    else:
                        logger.info('广西保监局 数据解析 ' + ' -- 数据已经存在')
                    result_map_list.append(each_map)
                if len(result_map_list) > 0:
                    logger.info('广西保监局 数据解析 ' + ' -- 一共有%d条数据' % len(result_map_list))
                    db.parsed_data.update_one({'_id': file_id}, {'$set': {'parsed': True}})
                    logger.info('广西保监局 数据解析 ' + ' -- 修改parsed完成')
                else:
                    logger.warning('广西保监局 数据解析 ' + ' -- 无数据')

        else:
            if '.doc' in content_text:
                if len([each_link.attrs['href'] for each_link in table_content.find_all('a')
                        if '.doc' in each_link.attrs['href'].strip()]) > 0:
                    doc_link = [urljoin(announcement_url, each_link.attrs['href'])
                                for each_link in table_content.find_all('a')
                                if '.doc' in each_link.attrs['href'].strip()][0]
                    link_type = doc_link.split('.')[-1]
                    response = request_site_page(doc_link)
                    if response is None:
                        logger.error('网页请求错误')
                        return

                    with open('./test/tmp.' + link_type, 'wb') as tmp_file:
                        for chunk in response.iter_content(chunk_size=1024):
                            if chunk:
                                tmp_file.write(chunk)

                    if not os.path.exists('./test/tmp.docx'):
                        shell_str = '/usr/local/bin/soffice --headless --convert-to docx ' + \
                                    './test/tmp.' + link_type + ' --outdir ./test'
                        process = subprocess.Popen(shell_str.replace(r'(', r'\(').replace(r')', r'\)'),
                                                   shell=True, stdout=subprocess.PIPE)
                        process.communicate()

                    with open('./test/tmp.docx', 'rb') as docx_file:
                        docx_content = docx_file.read()

                    if db.parsed_data.find({'origin_url': announcement_url,
                                            'oss_file_origin_url': doc_link}).count() == 0:
                        oss_file_map = {
                            'origin_url': announcement_url,
                            'oss_file_origin_url': doc_link,
                            'origin_url_id': each_circ_data['_id'],
                            'oss_file_type': 'docx',
                            'oss_file_name': announcement_title,
                            'oss_file_content': docx_content,
                            'parsed': False
                        }
                        response = db.parsed_data.insert_one(oss_file_map)
                        file_id = response.inserted_id
                        oss_add_file(ali_bucket, str(file_id) + '/' + announcement_title + '.docx', docx_content)
                        db.circ_data.update_one({'_id': each_circ_data['_id']}, {'$set': {'status': 'parsed'}})
                    else:
                        db.circ_data.update_one({'_id': each_circ_data['_id']}, {'$set': {'status': 'parsed'}})
                        file_id = db.parsed_data.find_one({'origin_url': announcement_url,
                                                           'oss_file_origin_url': doc_link})['_id']
                    document = Document('./test/tmp.docx')
                    logger.info('删除tmp文件')
                    if os.path.exists('./test/tmp.doc'):
                        os.remove('./test/tmp.doc')
                    if os.path.exists('./test/tmp.docx'):
                        os.remove('./test/tmp.docx')
                    result_map_list = []
                    tables = document.tables
                    for table in tables:
                        for row in table.rows:
                            result_map_list.append([kk.text for kk in row.cells])
                    parse_table(file_id, result_map_list, announcement_title, each_circ_data['publishDate'], db, logger)
            else:
                if db.parsed_data.find(
                        {'origin_url': announcement_url, 'oss_file_origin_url': announcement_url}).count() == 0:
                    oss_file_map = {
                        'origin_url': announcement_url,
                        'oss_file_origin_url': announcement_url,
                        'origin_url_id': each_circ_data['_id'],
                        'oss_file_type': 'html',
                        'oss_file_name': announcement_title,
                        'oss_file_content': r.text.encode(r.encoding).decode('utf-8'),
                        'parsed': False
                    }
                    insert_response = db.parsed_data.insert_one(oss_file_map)
                    file_id = insert_response.inserted_id
                    oss_add_file(ali_bucket, str(file_id) + '/' + announcement_title + '.html',
                                 r.text.encode(r.encoding).decode('utf-8'))
                    db.circ_data.update_one({'_id': each_circ_data['_id']}, {'$set': {'status': 'parsed'}})
                else:
                    db.circ_data.update_one({'_id': each_circ_data['_id']}, {'$set': {'status': 'parsed'}})
                    file_id = db.parsed_data.find_one({'origin_url': announcement_url,
                                                       'oss_file_origin_url': announcement_url})['_id']

                document_code_compiler = re.compile(r'(桂银?保监罚.\d{4}.\d+号)')
                if document_code_compiler.search(content_text):
                    document_code = document_code_compiler.search(content_text).group(1).strip()
                    litigant_compiler = re.compile(
                        document_code.replace(r'[', r'\[').replace(r']', r'\]') +
                        r'\n([\s\S]*?)\n' + r'(经查|经检查|依据.*?的有关规定|抽查|'
                                            r'经抽查|.*?(现场检查|案件调查))')
                    litigant = litigant_compiler.search(content_text).group(1).strip()
                else:
                    if document_code_compiler.search(title):
                        document_code = document_code_compiler.search(title).group(1).strip()
                    else:
                        document_code_compiler = re.compile(r'(.\d{4}.*?\d+号)')
                        if document_code_compiler.search(title):
                            document_code = '桂保监罚' + document_code_compiler.search(title).group(1).strip()
                        else:
                            document_code = ''
                    litigant_compiler = re.compile(r'^([\s\S]*?)\n' + r'(经查|经检查|依据.*?的有关规定|抽查|经抽查|'
                                                                      r'.*?(现场检查|案件调查))')
                    litigant = litigant_compiler.search(content_text).group(1).strip()

                truth_text_str = r'((经查)' \
                                 r'([\s\S]*?))' \
                                 r'((我局认为，)?(上述|以上).*?(事实|行为|事实).*?有.*?等证据(材料)?(在案)?证明(,|，|。)(足以认定。)?|' \
                                 r'(我局认为，|综上，)?(上述|以上).*?(行为|问题|事实).*?违反.*?第.*条.*?(的规定)?|' \
                                 r'上述违法事实有现场检查事实确认书、相关人员询问笔录、\n会计凭证复印件、营业执照复印件、劳动合同复印件、任职文件复印件等证据材料在案证明。|' \
                                 r'你单位的行为违反了《中华人民共和国保险法》第八十二条的规定|' \
                                 r'违反了《保险法》第一百零七条、第一百二十二条|' \
                                 r'(依据|根据)\n?.*?第)'
                truth_compiler = re.compile(truth_text_str)
                if truth_compiler.search(content_text):
                    truth = truth_compiler.search(content_text).group(1).strip()
                else:
                    truth_text_str = litigant.replace(r'(', r'\(').replace(r')', r'\)') \
                                     + r'([\s\S]*?)' \
                                       r'((我局认为，)?(上述|以上).*?(事实|行为|事实).?有.*?等证据(在案)?证明|' \
                                       r'(我局认为，|综上，)?(上述|以上).*?(行为|问题|事实).*?违反.*?第.*条.*?(的规定)?)'
                    truth_compiler = re.compile(truth_text_str)
                    truth = truth_compiler.search(content_text).group(1).strip()

                if '申辩' in content_text:
                    defense_text_str = r'((针对.*?行为.*?申辩意见|(当事人)?[^，。,；\n]*?(未)?提出(了)?陈述申辩(意见)?|' \
                                       r'[^，。,；\n]*?向我局(报送|递交|提出)[^，。,；\n]*?|本案在审理过程中.*?提出陈述申辩|' \
                                       r'[^，。,；\n]*?在(申辩材料|陈述申辩|陈述申辩意见|申辩意见)中称|[^，。,；\n]*?在听证阶段提出|' \
                                       r'[^，。,；\n]*?在法定期限内(未)?提出(了)?(听证要求|陈述申辩|陈述申辩及听证要求))' \
                                       r'([\s\S]*?))' \
                                       r'(因此，我局决定|' \
                                       r'我局经复核(认为|决定)|' \
                                       r'本案现已审理终结|' \
                                       r'我局经复查[^，。,；\n]*?情况|' \
                                       r'我局[^，。,；\n]*?认真复核|' \
                                       r'经研究，对[^，。,；\n]*?予以采纳。|' \
                                       r'我局认为.*?申辩理由|' \
                                       r'依据.*?我局认为.*?的申辩理由|' \
                                       r'经研究，我局认为.*?申辩意见|' \
                                       r'经我局审核，决定|' \
                                       r'我局认为，上述违法行为事实清楚、证据确凿、法律法规适当|' \
                                       r'我局对陈述申辩意见进行了复核|' \
                                       r'经我局审核|' \
                                       r'针对[^，。,；\n]*?的(陈述)?申辩意见，我局进行了核实|' \
                                       r'经查，我局认为|' \
                                       r'依据现场检查及听证情况|' \
                                       r'经查)'
                    defense_compiler = re.compile(defense_text_str, re.MULTILINE)
                    defense_list = defense_compiler.findall(content_text)
                    if len(defense_list) != 0:
                        defense = defense_list[-1][0].strip()
                        defense_response_str = defense.replace(r'[', r'\[').replace(r']', r'\]') \
                                               + r'(([\s\S]*?)' + r'(本案现已审理终结。|不符合.*?情形。|根据.*?依法可从轻或者减轻行政处罚。|' \
                                                                  r'对[^，。,；\n]*?申辩意见(不予|予以)采纳|因此.*?申辩理由.*?成立。|' \
                                                                  r'我局认为.*?申辩(理由|意见).*?符合.*?第.*?条.*?的条件.(予以采纳。)?))'
                        defense_response_compiler = re.compile(defense_response_str, re.MULTILINE)
                        if defense_response_compiler.search(content_text):
                            defense_response = defense_response_compiler.search(content_text).group(1).strip()
                        else:
                            if '未' in defense:
                                defense_response = ''
                    else:
                        defense_text_str = '([^。；\n]*?向.*?公告送达了《行政处罚事先告知书》.*?提出陈述申辩。|' \
                                           '我局依法于2012年5月25日对你公司送达了《行政处罚事先告知书》，你公司在规定的时间内未提出陈述和申辩意见，也未要求举行听证。)'
                        defense_compiler = re.compile(defense_text_str, re.MULTILINE)
                        defense = defense_compiler.search(content_text).group(1).strip()
                        defense_response = ''
                else:
                    defense = defense_response = ''

                punishment_decision_text_str = r'(((依据|根据)\n?.*?第?.*?条.*?(规定)?.?(我局)?(决定|责令|给予|于.*?向.*?发出|对.*?作出|拟对你)|' \
                                               r'我局决定.*?作出(如下|以下)(行政)?处罚：)' \
                                               r'([\s\S]*?))' \
                                               r'(请在本处罚决定书送达之日|当事人应当在接到本处罚决定书之日|如不服本处罚决定|' \
                                               r'请(在)?接到本处罚决定书之日|如不服从本处罚决定|如.*?对本处罚决定不服)'

                punishment_decision_compiler = re.compile(punishment_decision_text_str)
                punishment_decision = punishment_decision_compiler.search(content_text).group(1).strip()

                punishment_basis_str_list = [
                    r'([^\n。；]*?)(问题|行为|事项|情况|事实)([^\n。；\s]*?)违反.*?\n?.*?第.*?条?\n?.*?((的|之|等)(相关)?规定)?',
                    r'中国平安财产保险股份有限公司玉林中心支公司上述行为违反了《中华人民共和国保险法》，我局依法对该公司进行了行政处罚。曾家兴是对该公司上述违法行为负有直接责任的主管人员',
                    r'经查.*?存在.*?行为, 违反了.*?第.*?条的规定。\n.*?负有直接责任'
                ]
                punishment_basis_str = '|'.join(punishment_basis_str_list)
                punishment_basis_compiler = re.compile(r'[。\n；]' + '(' + punishment_basis_str + ')' +
                                                       '.(\n?依据|\n?根据|\n?鉴于|\n?你公司在申辩材料中称)', re.MULTILINE)
                punishment_basis_list = punishment_basis_compiler.findall(content_text)
                punishment_basis = '；'.join([kk[0].strip() for kk in punishment_basis_list])

                publish_date_text = re.search(
                    punishment_decision.replace(r'(', r'\(').replace(r')', r'\)').replace(r'[', r'\[').replace(r']',
                                                                                                               r'\]').
                    replace(r'*', r'\*') + r'([\s\S]*?)$', content_text).group(1).replace('\n', '')
                if re.search(r'.{4}年.{1,2}月.{1,3}日', publish_date_text):
                    publish_date = re.findall('.{4}年.{1,2}月.{1,3}日', publish_date_text)[-1].replace(' ', '')
                    m = re.match("([0-9零一二两三四五六七八九十〇○ＯOΟО]+年)?([0-9一二两三四五六七八九十]+)月?([0-9一二两三四五六七八九十]+)[号日]?",
                                 publish_date)
                    real_publish_date = get_year(m.group(1)) + str(cn2dig(m.group(2))) + '月' + str(
                        cn2dig(m.group(3))) + '日'
                else:
                    publish_date_text = table_content.find_all('tr')[1].text
                    publish_date = re.findall(r'\d{4}-\d{1,2}-\d{1,2}', publish_date_text)[-1]
                    real_publish_date = publish_date.split('-')[0] + '年' + str(
                        int(publish_date.split('-')[1])) + '月' + str(
                        int(publish_date.split('-')[2])) + '日'

                result_map = {
                    'announcementTitle': title,
                    'announcementOrg': '广西银保监局',
                    'announcementDate': real_publish_date,
                    'announcementCode': document_code,
                    'facts': truth,
                    'defenseOpinion': defense,
                    'defenseResponse': defense_response,
                    'litigant': litigant[:-1] if litigant[-1] == '：' else litigant,
                    'punishmentBasement': punishment_basis,
                    'punishmentDecision': punishment_decision,
                    'type': '行政处罚决定',
                    'oss_file_id': file_id,
                    'status': 'not checked'
                }
                logger.info(result_map)
                if db.announcement.find({'announcementTitle': title, 'oss_file_id': file_id}).count() == 0:
                    db.announcement.insert_one(result_map)
                    logger.info('广西保监局 数据解析 ' + ' -- 数据导入完成')
                else:
                    logger.info('广西保监局 数据解析 ' + ' -- 数据已经存在')
                db.parsed_data.update_one({'_id': file_id}, {'$set': {'parsed': True}})
                logger.info('广西保监局 数据解析 ' + ' -- 修改parsed完成')
